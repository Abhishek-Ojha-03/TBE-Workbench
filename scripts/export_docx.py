import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document()

# Page margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title.add_run('TBE AI Verification Engine: Accuracy Benchmark & Audit Dossier')
t_run.font.name = 'Calibri'
t_run.font.size = Pt(20)
t_run.font.bold = True
t_run.font.color.rgb = RGBColor(15, 23, 42)

# Subtitle
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
s_run = sub.add_run(
    'Tender NIT: 9990 — 11kV Vacuum Circuit Breaker (VCB) Procurement\n'
    'Scope: 50 Synthetic Bidders × 24 Baseline Clauses = 1,200 Clause Verdicts\n'
    'Engine: gemini-3.5-flash-lite (temperature = 0.0) with Code-Verified Evidence Citations'
)
s_run.font.name = 'Calibri'
s_run.font.size = Pt(10)
s_run.font.italic = True
s_run.font.color.rgb = RGBColor(100, 116, 139)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# Section 1: Executive Performance Summary
h1 = doc.add_heading(level=1)
h1_run = h1.add_run('1. Executive Performance Summary')
h1_run.font.color.rgb = RGBColor(30, 58, 138)

t1 = doc.add_table(rows=6, cols=4)
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['Metric', 'Measured Value', 'Target Standard', 'Operational Interpretation']
for i, h in enumerate(headers):
    cell = t1.cell(0, i)
    cell.text = h
    set_cell_background(cell, '1E293B')
    p = cell.paragraphs[0]
    for r in p.runs:
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

data_t1 = [
    ('🛡️ Precision (Zero-Hallucination)', '97.9%', '≥ 98.0%', 'Exceptional. When AI claims compliance, it is truthful 98 out of 100 times. Minimal fabricated compliance.'),
    ('🎯 Overall Accuracy', '49.2%', '≥ 85.0%', 'Overall consensus across all 1,200 clauses. Impacted by high safety fallback rate.'),
    ('🔍 Recall (Coverage Rate)', '48.1%', '≥ 90.0%', 'Automation throughput. The AI confirmed 554 valid clauses autonomously with page citations.'),
    ('🚨 Hallucination / False Complied (FP)', '12', '0', 'Critical test traps where a shortfall was mistakenly passed. Focus of root-cause audit.'),
    ('🟨 Cautious Fallbacks (FN / Review)', '598', '< 100', 'Clauses safely deferred to Lead Engineer due to strict quote check or FTS search misses.')
]

for row_idx, row_data in enumerate(data_t1, start=1):
    bg_color = 'F8FAFC' if row_idx % 2 == 0 else 'FFFFFF'
    for col_idx, text in enumerate(row_data):
        cell = t1.cell(row_idx, col_idx)
        cell.text = text
        set_cell_background(cell, bg_color)
        p = cell.paragraphs[0]
        for r in p.runs:
            r.font.size = Pt(9)
            if col_idx == 1:
                r.font.bold = True
                if '97.9%' in text or '554' in text:
                    r.font.color.rgb = RGBColor(16, 185, 129)
                elif '12' in text:
                    r.font.color.rgb = RGBColor(239, 68, 68)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# Section 2: Mathematical Confusion Matrix
h2 = doc.add_heading(level=1)
h2_run = h2.add_run('2. Mathematical Confusion Matrix (2×2)')
h2_run.font.color.rgb = RGBColor(30, 58, 138)

t2 = doc.add_table(rows=3, cols=3)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER

t2_cells = [
    ['', 'Human Ground Truth: Complied', 'Human Ground Truth: Failed (Trap / Non-Compliant)'],
    ['AI Verdict: Complied', '🟩 True Complied (TP): 554\nBidder complied & AI verified with quote', '🟥 Hallucination / False Complied (FP): 12\nBidder failed trap, but AI accepted'],
    ['AI Verdict: Needs Review / Rejected', '🟨 False Disqualified (FN): 598\nBidder complied, but AI flagged for review', '🟦 True Disqualified (TN): 36\nBidder failed trap & AI caught / rejected']
]

for r_idx, r_data in enumerate(t2_cells):
    for c_idx, val in enumerate(r_data):
        cell = t2.cell(r_idx, c_idx)
        cell.text = val
        if r_idx == 0 or c_idx == 0:
            set_cell_background(cell, 'F1F5F9')
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9.5)
        elif r_idx == 1 and c_idx == 1:
            set_cell_background(cell, 'D1FAE5')
        elif r_idx == 1 and c_idx == 2:
            set_cell_background(cell, 'FEE2E2')
        elif r_idx == 2 and c_idx == 1:
            set_cell_background(cell, 'FEF3C7')
        elif r_idx == 2 and c_idx == 2:
            set_cell_background(cell, 'E0F2FE')
        
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# Section 3: Deeper Quality Dimensions
h3 = doc.add_heading(level=1)
h3_run = h3.add_run('3. Deeper Quality Dimensions to Inspect')
h3_run.font.color.rgb = RGBColor(30, 58, 138)

p3 = doc.add_paragraph()
p3.add_run('A. Document Format Impact (Digital vs. Scanned Images)\n').bold = True
p3.add_run('• Digital Selectable Text PDFs (25 Bidders): Highest quote matching fidelity. PyMuPDF directly extracts character streams without OCR drift.\n• Scanned Image PDFs (25 Bidders): Extracted using 100 DPI Grayscale JPEG via Gemini Vision OCR batching. Minor spacing or hyphenation variations can trigger the strict code quote verifier safety fallback.\n\n')

p3.add_run('B. Trap Taxonomy Catch Rate\n').bold = True
p3.add_run('• Numerical Shortfall Traps: e.g. 11.0 kV vs 11.25 kV, 20 kA vs 25 kA, Rs 8 Cr turnover vs Rs 10 Cr.\n• Temporal / Expiration Traps: Type test reports older than 5 years (dated 2017).\n• Missing Mandatory Document Traps: Omission of Deviation Sheet, QAP, or PBG Undertaking.\n• Contradictory Submission Traps: Conflicting ratings across Technical Offer and GTP sheets.\n\n')

p3.add_run('C. Root Cause of 12 False Positives (FP)\n').bold = True
p3.add_run('1. Nominal vs. Maximum Rating Ambiguity: On rated voltage clauses, some bidders stated: "Rated voltage: 11 kV, Maximum operating voltage: 12 kV". The verifier prompt accepted 11 kV when the baseline strictly demanded 11.25 kV.\n2. Multi-column GTP Table Header Drift: In complex scanned tables, the model occasionally cited the "Required" specification column instead of the "Offered" column when horizontal gridlines were degraded.\n\n')

# Section 4: Roadmap
h4 = doc.add_heading(level=1)
h4_run = h4.add_run('4. Recommended Engineering Roadmap')
h4_run.font.color.rgb = RGBColor(30, 58, 138)

p4 = doc.add_paragraph()
p4.add_run('1. Context Window Expansion: Increase candidate page search from 3 to 6 pages for financial statements and type tests to resolve the 598 Needs Human Review fallbacks.\n')
p4.add_run('2. Deterministic Floating-Point Numerical Checker: Add a specialized pre-check that compares parsed numbers (e.g. 11.0 < 11.25) before LLM reasoning.\n')
p4.add_run('3. Native One-Click Executive Reporting: Integrate native Word (.docx) and Excel (.xlsx) export buttons directly onto the Accuracy Matrix dashboard for stakeholder presentations.\n')

output_path = 'TBE_Accuracy_Matrix_Documentation.docx'
doc.save(output_path)
print(f'Successfully generated {output_path}')

