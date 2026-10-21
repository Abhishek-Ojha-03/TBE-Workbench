"""
Unit test for Step 6: Create the working Excel.
Done when: The generated file looks the same as the template when you open both side by side.
Verifies:
1. Populates template without creating from scratch.
2. Structure: QR rows first, then ONE BLANK ROW, then other requirements.
3. Accurate column groups per bidder.
4. Cell coordinate map accurately links (requirement_id, bidder_id) to (row, col).
"""
import os
import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from db.models import Base, Tender, BaselineVersion, Requirement, Bidder
from outputs.excel_generator import WorkingExcelGenerator

TEST_STEP6_DB = "test_step6.db"
TEST_URL = f"sqlite:///{TEST_STEP6_DB}"


@pytest.fixture(scope="module")
def step6_db():
    engine = create_engine(TEST_URL, connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = Session()
    yield db
    db.close()
    engine.dispose()
    if os.path.exists(TEST_STEP6_DB):
        try:
            os.remove(TEST_STEP6_DB)
        except PermissionError:
            pass


def test_working_excel_structure(step6_db):
    # 1. Setup Tender
    tender = Tender(description="Supply of 11kV Switchgear", nit_no="NIT-STEP6-VCB", status="open")
    step6_db.add(tender)
    step6_db.commit()

    # 2. Add Baseline Version 1
    v1 = BaselineVersion(tender_id=tender.id, version_number=1)
    step6_db.add(v1)
    step6_db.commit()

    # 2 QR requirements and 2 Technical requirements
    r1 = Requirement(baseline_version_id=v1.id, clause="QR-1", text="5 yrs OEM Experience", is_QR=True)
    r2 = Requirement(baseline_version_id=v1.id, clause="QR-2", text="Supplied 500 panels", is_QR=True)
    r3 = Requirement(baseline_version_id=v1.id, clause="4.1", text="11kV rated voltage", is_QR=False)
    r4 = Requirement(baseline_version_id=v1.id, clause="4.2", text="25kA short circuit rating", is_QR=False)
    step6_db.add_all([r1, r2, r3, r4])

    # 3. Add 2 Bidders
    b1 = Bidder(tender_id=tender.id, name="ABB India")
    b2 = Bidder(tender_id=tender.id, name="Siemens India")
    step6_db.add_all([b1, b2])
    step6_db.commit()

    # 4. Generate Working Excel
    gen = WorkingExcelGenerator()
    excel_path, cell_map = gen.generate_working_excel(step6_db, tender, output_filename="test_step6_output.xlsx")

    assert os.path.exists(excel_path)
    assert len(cell_map) == 8  # 4 requirements x 2 bidders

    # 5. Open and inspect Excel layout
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    # Check Tender Metadata
    assert ws["B2"].value == "NIT-STEP6-VCB"
    assert ws["B3"].value == "Supply of 11kV Switchgear"

    # Check Bidder Headers on Row 5
    assert "ABB INDIA" in str(ws["F5"].value)
    assert "SIEMENS INDIA" in str(ws["I5"].value)

    # Check Row sequence:
    # Row 7 = QR-1 (First QR row)
    assert ws["B7"].value == "Qualification (QR)"
    assert ws["C7"].value == "QR-1"

    # Row 8 = QR-2 (Second QR row)
    assert ws["B8"].value == "Qualification (QR)"
    assert ws["C8"].value == "QR-2"

    # Row 9 = BLANK SEPARATOR ROW
    assert ws["A9"].value is None
    assert ws["B9"].value is None

    # Row 10 = First other requirement (Clause 4.1)
    assert ws["B10"].value == "Technical Spec"
    assert ws["C10"].value == "4.1"

    # Row 11 = Second other requirement (Clause 4.2)
    assert ws["B11"].value == "Technical Spec"
    assert ws["C11"].value == "4.2"

    # Verify cell map coordinates
    assert cell_map[(r1.id, b1.id)] == (7, 7)  # Row 7, Col G (Compliance for Bidder 1)
    assert cell_map[(r1.id, b2.id)] == (7, 10) # Row 7, Col J (Compliance for Bidder 2)
    assert cell_map[(r3.id, b1.id)] == (10, 7) # Row 10, Col G (After blank row)

    wb.close()
    if os.path.exists(excel_path):
        os.remove(excel_path)

